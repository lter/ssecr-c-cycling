"""Assemble the Supplementary Information from pipeline outputs.

Writes docs/supplement/Supplementary_Information.docx (Fig. S1 and the table captions)
and docs/supplement/Supplementary_Tables_S1-S3.xlsx (the three tables, one sheet each). Run from the repo root after analysis/08 and 99.
"""
import csv
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

SUP = "figures/supplemental"
TABLES = {
    "Table S1": (f"{SUP}/TableS1_full_treatment_statistics.csv",
                 "Per-treatment statistics for the 83 classified treatment-experiment combinations: trend class, "
                 "slope, P-value and R-squared of the linear fit of the response ratio over time, coefficient of "
                 "variation (CV), mean response ratio, number of timepoints, record length, early (first three "
                 "calendar years) and full-duration log response ratios (LRR) with their difference and "
                 "classification (accurate, underestimate, overestimate, wrong direction; tolerance 0.2), and years "
                 "to first detection of a significant trend (directional treatments only)."),
    "Table S2": (f"{SUP}/TableS2_sizer_results.csv",
                 "SiZer results for the 79 treatment-experiment combinations with at least five timepoints: number "
                 "of significant slope changes at a 5-year bandwidth, the years at which they occurred, and segment "
                 "slopes with P-values."),
    "Table S3": (f"{SUP}/TableS3_case_study_details.csv",
                 "Details of the 18 case-study datasets: setting, experiment, manipulation, response variable and "
                 "type, years analyzed, number of years sampled, number of analyzable treatments (in parentheses, "
                 "the number screened), EDI data-package identifier, and data citation."),
}
FIG_S1 = ("figures/FigS1_early_vs_full_sensitivity.png",
          open("docs/figure_captions.txt", encoding="utf-8").read().split("FigS1_early_vs_full_sensitivity")[1]
          .split("-\n", 1)[1].strip().replace("-\n", "-").replace("\n", " "))

def read(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))

def num(v):
    try:
        return float(v)
    except ValueError:
        return v

# ---- workbook
wb = Workbook(); wb.remove(wb.active)
for name, (path, cap) in TABLES.items():
    ws = wb.create_sheet(name); rows = read(path)
    ws.append([f"{name}. {cap}"]); ws.append([])
    for r in rows:
        ws.append([num(v) for v in r])
    for c in ws[3]: c.font = Font(bold=True)
    ws["A1"].alignment = Alignment(wrap_text=False); ws.freeze_panes = "A4"
    for i, h in enumerate(rows[0], 1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = min(60, max(10, len(h) + 2))
wb.save("docs/supplement/Supplementary_Tables_S1-S3.xlsx")

# ---- document
d = Document()
st = d.styles["Normal"]; st.font.name = "Cambria"; st.font.size = Pt(11)
d.add_heading("Supplementary Information", level=1)
d.add_paragraph("Figure S1 and Tables S1–S3. The tables are provided in the accompanying workbook "
                "(Supplementary_Tables_S1-S3.xlsx), one sheet per table; their captions are given below.")
d.add_picture(FIG_S1[0], width=Inches(6.0))
p = d.add_paragraph(); p.add_run("Figure S1. ").bold = True; p.add_run(FIG_S1[1])
for name in TABLES:
    p = d.add_paragraph(); p.add_run(f"{name}. ").bold = True; p.add_run(TABLES[name][1] + " (Workbook.)")
d.save("docs/supplement/Supplementary_Information.docx")
print("written")
